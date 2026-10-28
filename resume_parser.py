from io import BytesIO

from docx import Document
from pypdf import PdfReader


def extract_resume_text(uploaded_file):
    """Read a PDF or DOCX resume without saving it to disk."""
    file_data = uploaded_file.getvalue()

    if len(file_data) > 10 * 1024 * 1024:
        raise ValueError("Please upload a resume smaller than 10 MB.")

    filename = uploaded_file.name.lower()

    try:
        if filename.endswith(".pdf"):
            reader = PdfReader(BytesIO(file_data))

            if reader.is_encrypted:
                raise ValueError(
                    "This PDF is password protected. "
                    "Please upload an unlocked copy."
                )

            text = "\n\n".join(
                page.extract_text() or ""
                for page in reader.pages
            )

        elif filename.endswith(".docx"):
            document = Document(BytesIO(file_data))
            sections = [
                paragraph.text
                for paragraph in document.paragraphs
                if paragraph.text.strip()
            ]

            # Include content stored in Word tables.
            for table in document.tables:
                for row in table.rows:
                    sections.append(
                        " | ".join(cell.text for cell in row.cells)
                    )

            text = "\n".join(sections)

        else:
            raise ValueError("Please upload a PDF or DOCX file.")

    except ValueError:
        raise
    except Exception as error:
        raise ValueError(
            "We couldn't read this file. "
            "Please try another PDF or DOCX copy."
        ) from error

    text = text.strip()

    if not text:
        raise ValueError(
            "No readable text was found. If your resume is a scanned "
            "image, please use a text-based PDF or paste its text."
        )

    return text